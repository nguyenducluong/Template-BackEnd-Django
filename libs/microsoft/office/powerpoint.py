"""Doc file PowerPoint (.pptx) theo path - 2 backend: python-pptx va COM (PowerPoint that).

BACKEND
-------
1. python-pptx (mac dinh) - thuan Python, chay duoc headless / Linux / Docker.
2. win32com (fallback) - chi khi python-pptx doc loi (VD .ppt cu, file mac dinh
   khac) VA may dang cai MS PowerPoint. Can ``pywin32``.

Ket qua gom danh sach ``slides``; moi slide co ``texts`` (o co chu), ``notes``
(ghi chu) va ``index``.
"""

import logging

from libs.microsoft.office._common import has_com, require_file, run_com, to_text
from libs.microsoft.office.errors import (
    PowerPointBackendUnavailable,
    PowerPointFileNotFound,
)

logger = logging.getLogger("apps")


# --- backend 1: python-pptx (mac dinh) ---


def _read_pptx_python(file_path, with_notes=True):
    """Doc .pptx bang python-pptx -> dict {slides: [...]}."""
    from pptx import Presentation

    presentation = Presentation(file_path)
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = to_text(shape.text_frame.text)
                if text:
                    texts.append(text)
        notes = ""
        if with_notes and getattr(slide, "has_notes_slide", False):
            notes = to_text(slide.notes_slide.notes_text_frame.text)
        slides.append({"index": index, "texts": texts, "notes": notes, "text": "\n".join(texts)})
    return {"slides": slides}


# --- backend 2: win32com (fallback) ---


def _read_pptx_com(file_path, with_notes=True):
    """Doc .pptx bang COM (PowerPoint that) -> dict {slides: [...]}."""

    def work(app):
        presentation = app.Presentations.Open(file_path, ReadOnly=True, WithWindow=False)
        try:
            slides = []
            for index in range(1, presentation.Slides.Count + 1):
                slide = presentation.Slides.Item(index)
                texts = []
                for shape in slide.Shapes:
                    try:
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            text = to_text(shape.TextFrame.TextRange.Text)
                            if text:
                                texts.append(text)
                    except Exception:  # pragma: no cover - shape khong ho tro text
                        continue
                notes = ""
                if with_notes:
                    try:
                        notes = to_text(slide.NotesPage.Shapes.Item(2).TextFrame.TextRange.Text)
                    except Exception:  # pragma: no cover - slide khong co ghi chu
                        notes = ""
                slides.append({"index": index, "texts": texts, "notes": notes, "text": "\n".join(texts)})
            return {"slides": slides}
        finally:
            presentation.Close()

    return run_com("PowerPoint.Application", work)


# --- API cong khai ---


def read_powerpoint(file_path, with_notes=True):
    """Doc file PowerPoint va tra ve du lieu JSON-serializable.

    Args:
        file_path (str): duong dan file .pptx (hoac .ppt neu co MS PowerPoint).
        with_notes (bool): ``True`` -> trich xuat ca ghi chu thuyet trinh.

    Returns:
        dict: ``{"slides": [{"index": 1, "texts": [...], "notes": str, "text": str}, ...]}``

    Raises:
        PowerPointFileNotFound: file khong ton tai.
        PowerPointBackendUnavailable: khong backend nao doc duoc.

    Vi du:
        >>> deck = read_powerpoint("D:/data/bai-giao-trinh.pptx")
        >>> len(deck["slides"])
    """
    require_file(file_path, PowerPointFileNotFound)

    try:
        return _read_pptx_python(file_path, with_notes)
    except Exception as python_error:
        logger.warning(
            "python-pptx failed to read %s (%s) - trying COM fallback",
            file_path,
            python_error,
        )
        if not has_com():
            raise PowerPointBackendUnavailable(
                "Cannot read PowerPoint file: {} (python-pptx failed: {}; "
                "pywin32 not available so COM fallback is skipped - "
                ".ppt needs MS PowerPoint on Windows)".format(file_path, python_error)
            ) from python_error
        try:
            return _read_pptx_com(file_path, with_notes)
        except PowerPointFileNotFound:
            raise
        except Exception as com_error:
            raise PowerPointBackendUnavailable(
                "Cannot read PowerPoint file: {} (python-pptx: {}; COM: {})".format(file_path, python_error, com_error)
            ) from com_error